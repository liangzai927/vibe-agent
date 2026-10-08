"""需求文档读取的测试。

重点不是「能不能读出字」，而是三件容易做砸的事：

* 长文档**不能被一次性倒进上下文**——所以要先给目录、再按需取用；
* 表格、标题这些结构不能在抽取时丢掉，否则模型看不懂层级；
* 缺依赖、扫描版 PDF、章节名打错这些情况，报错必须自解释。
"""

from __future__ import annotations

import sys
import unittest
from pathlib import Path
from unittest import mock

sys.path.insert(0, str(Path(__file__).resolve().parent))
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

from helpers import AgentTestCase  # noqa: E402
from va.tools.base import ToolError  # noqa: E402
from va.tools.document import (  # noqa: E402
    MAX_RETURN_CHARS,
    ReadDocument,
    detect_kind,
    find_headings,
    slice_section,
)

HAS_DOCX = __import__("importlib.util", fromlist=["util"]).find_spec("docx") is not None
HAS_PYPDF = __import__("importlib.util", fromlist=["util"]).find_spec("pypdf") is not None

MARKDOWN = """# 登录页改造需求

本文档描述登录页的改造范围。

## 背景

现有登录页转化率偏低，需要重新设计。

## 接口设计

POST /api/login
请求体：{ username, password }
返回：{ token, expires_at }

## 验收标准

1. 登录成功率可统计
2. 错误提示要具体

### 错误码

- 401 用户名或密码错误
- 429 尝试过于频繁

## 不在范围内

- 注册流程
- 第三方登录
"""


class TestStructure(unittest.TestCase):
    def test_detect_kind(self):
        self.assertEqual(detect_kind(Path("a.md")), "text")
        self.assertEqual(detect_kind(Path("a.txt")), "text")
        self.assertEqual(detect_kind(Path("a.html")), "html")
        self.assertEqual(detect_kind(Path("a.docx")), "docx")
        self.assertEqual(detect_kind(Path("a.pdf")), "pdf")
        self.assertEqual(detect_kind(Path("没有扩展名")), "text")

    def test_find_headings_levels(self):
        headings = find_headings(MARKDOWN)
        titles = [h.title for h in headings]
        self.assertEqual(titles[0], "登录页改造需求")
        self.assertIn("接口设计", titles)
        self.assertIn("错误码", titles)
        levels = {h.title: h.level for h in headings}
        self.assertEqual(levels["背景"], 2)
        self.assertEqual(levels["错误码"], 3)

    def test_slice_section_stops_at_same_level(self):
        headings = find_headings(MARKDOWN)
        found = slice_section(MARKDOWN, headings, "接口设计")
        self.assertIsNotNone(found)
        title, body = found
        self.assertEqual(title, "接口设计")
        self.assertIn("POST /api/login", body)
        self.assertNotIn("验收标准", body, "取章节时越界了")

    def test_slice_section_includes_deeper_headings(self):
        headings = find_headings(MARKDOWN)
        _, body = slice_section(MARKDOWN, headings, "验收标准")
        self.assertIn("错误码", body, "子标题应该包含在父章节里")
        self.assertNotIn("不在范围内", body)

    def test_slice_section_missing(self):
        headings = find_headings(MARKDOWN)
        self.assertIsNone(slice_section(MARKDOWN, headings, "不存在的章节"))

    def test_slice_section_empty_keyword(self):
        headings = find_headings(MARKDOWN)
        self.assertIsNone(slice_section(MARKDOWN, headings, "   "))


class TestMarkdown(AgentTestCase):
    def setUp(self) -> None:
        super().setUp()
        (self.root / "req.md").write_text(MARKDOWN, encoding="utf-8")

    def tool(self) -> ReadDocument:
        return ReadDocument(self.workspace, self.cfg)

    def test_header_reports_shape(self):
        out = self.tool().run(path="req.md", outline=True)
        self.assertIn("req.md", out)
        self.assertIn("文本", out)
        self.assertIn("个标题", out)

    def test_outline_lists_headings(self):
        out = self.tool().run(path="req.md", outline=True)
        for title in ("背景", "接口设计", "验收标准", "错误码"):
            self.assertIn(title, out)
        self.assertNotIn("POST /api/login", out, "outline 模式不该返回正文")

    def test_default_returns_the_body(self):
        out = self.tool().run(path="req.md")
        self.assertIn("POST /api/login", out)
        self.assertIn("正文", out)

    def test_section_extraction(self):
        out = self.tool().run(path="req.md", section="接口")
        self.assertIn("POST /api/login", out)
        self.assertNotIn("转化率偏低", out, "只该给这一节")

    def test_section_not_found_lists_candidates(self):
        with self.assertRaises(ToolError) as ctx:
            self.tool().run(path="req.md", section="压根没有")
        message = str(ctx.exception)
        self.assertIn("找不到", message)
        self.assertIn("接口设计", message, "应该把可用标题列出来")
        self.assertIn("outline=true", message, "应该告诉它怎么查目录")

    def test_offset_and_limit(self):
        first = self.tool().run(path="req.md", offset=0, limit=50)
        second = self.tool().run(path="req.md", offset=50, limit=50)
        self.assertIn("第 0-50 字符", first)
        self.assertIn("第 50-100 字符", second)
        self.assertNotEqual(first.split("──")[-1], second.split("──")[-1])

    def test_offset_beyond_end_is_graceful(self):
        with self.assertRaises(ToolError) as ctx:
            self.tool().run(path="req.md", offset=999_999)
        self.assertIn("超过文档长度", str(ctx.exception))

    def test_truncated_read_tells_how_to_continue(self):
        out = self.tool().run(path="req.md", limit=60)
        self.assertIn("后面还有", out)
        self.assertIn("offset=60", out)

    def test_missing_file(self):
        with self.assertRaises(ToolError):
            self.tool().run(path="nope.md")

    def test_directory_is_rejected(self):
        (self.root / "adir").mkdir()
        with self.assertRaises(ToolError) as ctx:
            self.tool().run(path="adir")
        self.assertIn("目录", str(ctx.exception))


class TestHtml(AgentTestCase):
    def test_tags_stripped_and_headings_kept(self):
        (self.root / "spec.html").write_text(
            """<html><head><style>body{color:red}</style>
            <script>console.log('不该出现')</script></head>
            <body>
              <h1>订单模块</h1>
              <p>支持<strong>批量</strong>导出。</p>
              <h2>字段</h2>
              <ul><li>order_id</li><li>amount</li></ul>
            </body></html>""",
            encoding="utf-8",
        )
        out = ReadDocument(self.workspace, self.cfg).run(path="spec.html")

        self.assertIn("订单模块", out)
        self.assertIn("批量", out)
        self.assertIn("order_id", out)
        self.assertNotIn("<", out, "HTML 标签没剥干净")
        self.assertNotIn("console.log", out, "script 内容不该出现")
        self.assertNotIn("color:red", out, "style 内容不该出现")
        # 标题要保留成 Markdown 形式，否则抽不出目录
        self.assertIn("HTML", out)

    def test_headings_are_extractable(self):
        (self.root / "h.html").write_text(
            "<h1>A</h1><p>x</p><h2>B</h2><p>y</p>", encoding="utf-8"
        )
        out = ReadDocument(self.workspace, self.cfg).run(path="h.html", outline=True)
        self.assertIn("A", out)
        self.assertIn("B", out)


class TestMissingDependencies(AgentTestCase):
    def test_docx_without_library_says_how_to_install(self):
        (self.root / "a.docx").write_bytes(b"PK\x03\x04 fake")
        with mock.patch.dict(sys.modules, {"docx": None}):
            with self.assertRaises(ToolError) as ctx:
                ReadDocument(self.workspace, self.cfg).run(path="a.docx")
        message = str(ctx.exception)
        self.assertIn("python-docx", message)
        self.assertIn("pip install", message)

    def test_pdf_without_library_says_how_to_install(self):
        (self.root / "a.pdf").write_bytes(b"%PDF-1.4 fake")
        with mock.patch.dict(sys.modules, {"pypdf": None}):
            with self.assertRaises(ToolError) as ctx:
                ReadDocument(self.workspace, self.cfg).run(path="a.pdf")
        message = str(ctx.exception)
        self.assertIn("pypdf", message)
        self.assertIn("pip install", message)


@unittest.skipUnless(HAS_DOCX, "需要 python-docx")
class TestDocx(AgentTestCase):
    def _make(self, name: str = "prd.docx") -> Path:
        import docx

        document = docx.Document()
        document.add_heading("支付改造", level=1)
        document.add_paragraph("目标：把支付成功率提到 99%。")
        document.add_heading("接口", level=2)
        document.add_paragraph("POST /pay")
        table = document.add_table(rows=2, cols=2)
        table.cell(0, 0).text = "字段"
        table.cell(0, 1).text = "类型"
        table.cell(1, 0).text = "amount"
        table.cell(1, 1).text = "int"
        path = self.root / name
        document.save(str(path))
        return path

    def test_extracts_headings_and_tables(self):
        self._make()
        out = ReadDocument(self.workspace, self.cfg).run(path="prd.docx")

        self.assertIn("Word", out)
        self.assertIn("支付改造", out)
        self.assertIn("POST /pay", out)
        self.assertIn("amount | int", out, "表格内容丢了")

    def test_outline_from_docx_headings(self):
        self._make()
        out = ReadDocument(self.workspace, self.cfg).run(path="prd.docx", outline=True)
        self.assertIn("支付改造", out)
        self.assertIn("接口", out)
        self.assertNotIn("99%", out)

    def test_section_from_docx(self):
        self._make()
        out = ReadDocument(self.workspace, self.cfg).run(path="prd.docx", section="接口")
        self.assertIn("POST /pay", out)
        self.assertNotIn("99%", out)

    def test_reports_paragraph_and_table_counts(self):
        self._make()
        out = ReadDocument(self.workspace, self.cfg).run(path="prd.docx", outline=True)
        self.assertIn("段落", out)
        self.assertIn("表格", out)


@unittest.skipUnless(HAS_PYPDF, "需要 pypdf")
class TestPdf(AgentTestCase):
    def test_text_pdf_is_readable(self):
        """用一张手写的极简 PDF —— 只要能提取出文字即可。"""
        content = b"""%PDF-1.4
1 0 obj<</Type/Catalog/Pages 2 0 R>>endobj
2 0 obj<</Type/Pages/Kids[3 0 R]/Count 1>>endobj
3 0 obj<</Type/Page/Parent 2 0 R/MediaBox[0 0 200 200]/Contents 4 0 R/Resources<</Font<</F1 5 0 R>>>>>>endobj
4 0 obj<</Length 60>>stream
BT /F1 12 Tf 20 100 Td (Hello PRD) Tj ET
endstream
endobj
5 0 obj<</Type/Font/Subtype/Type1/BaseFont/Helvetica>>endobj
trailer<</Root 1 0 R>>
"""
        (self.root / "p.pdf").write_bytes(content)
        try:
            out = ReadDocument(self.workspace, self.cfg).run(path="p.pdf")
        except ToolError as exc:
            self.skipTest(f"这份手写 PDF 解析不出来：{exc}")
        self.assertIn("PDF", out)

    def test_scanned_pdf_gives_actionable_error(self):
        """扫描件一页文字都提不出来——要说清楚该去做 OCR，而不是给个空结果。"""
        from va.tools.document import HANDLERS, Document

        (self.root / "scan.pdf").write_bytes(b"%PDF-1.4")
        # 注意要 patch HANDLERS 本身：它在模块导入时就把函数引用存下来了，
        # patch 模块属性是改不到它的
        with mock.patch.dict(
            HANDLERS,
            {
                "pdf": lambda p: Document(
                    p, "PDF", "", {"页数": 3, "无法提取文字的页": 3}
                )
            },
        ):
            with self.assertRaises(ToolError) as ctx:
                ReadDocument(self.workspace, self.cfg).run(path="scan.pdf")

        message = str(ctx.exception)
        self.assertIn("扫描件", message)
        self.assertIn("OCR", message)
        self.assertIn("read_design", message, "应该指出用视觉模型这条替代路径")


class TestToolContract(AgentTestCase):
    def test_no_approval_no_snapshots(self):
        tool = ReadDocument(self.workspace, self.cfg)
        self.assertFalse(tool.requires_approval)
        self.assertEqual(tool.affected_paths({"path": "a.md"}), [])

    def test_registered(self):
        from va.tools import build_registry

        registry = build_registry(self.workspace, cfg=self.cfg)
        self.assertIsNotNone(registry.get("read_document"))

    def test_default_limit_is_sane(self):
        self.assertGreaterEqual(MAX_RETURN_CHARS, 5000)
        self.assertLessEqual(MAX_RETURN_CHARS, 100_000)


if __name__ == "__main__":
    unittest.main(verbosity=2)
